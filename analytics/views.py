from django.shortcuts import get_object_or_404
from django.db.models import Count, Avg, Q, F, ExpressionWrapper, fields, Sum
from django.db.models.functions import TruncDate, ExtractDay, TruncMonth
from django.utils.timezone import now, timedelta
from rest_framework import generics, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView
from django_filters.rest_framework import DjangoFilterBackend
from django_filters import FilterSet, DateFilter
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import TemplateView
from django.http import HttpResponse
import openpyxl
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from io import BytesIO
from apps.organization.models import EmployeeProfile, ManagerProfile
from apps.tasks.models import Task, TimeLog
from apps.submissions.models import Submission
from .models import KPI, Report, AnalyticsCache
from .serializers import (
    ProductivitySerializer, DepartmentAnalyticsSerializer,
    TaskTrendsSerializer, ReportSerializer
)
from .permissions import IsAdminOrManager, DepartmentAccessPermission, EmployeeAnalyticsPermission
from .forms import DateRangeForm

class ProductivityFilter(FilterSet):
    start_date = DateFilter(field_name='created_at', lookup_expr='gte')
    end_date = DateFilter(field_name='created_at', lookup_expr='lte')

    class Meta:
        model = Task
        fields = ['start_date', 'end_date']

class EmployeeProductivityView(generics.ListAPIView):
    serializer_class = ProductivitySerializer
    permission_classes = [IsAuthenticated, IsAdminOrManager]
    filter_backends = [DjangoFilterBackend]
    filterset_class = ProductivityFilter

    def get_queryset(self):
        # This view computes productivity metrics on the fly
        # For scalability, consider caching or pre-computing
        start_date = self.request.query_params.get('start_date')
        end_date = self.request.query_params.get('end_date')

        tasks = Task.objects.select_related('assigned_to').prefetch_related('submissions')

        if start_date and end_date:
            tasks = tasks.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

        # Aggregate productivity data
        productivity_data = []
        users = set()
        for task in tasks:
            users.add(task.assigned_to)

        for user in users:
            user_tasks = tasks.filter(assigned_to=user)
            completed_tasks = user_tasks.filter(status='COMPLETED').count()
            overdue_tasks = user_tasks.filter(status='OVERDUE').count()
            total_tasks = user_tasks.count()

            # Calculate average completion time (simplified)
            completed_submissions = Submission.objects.filter(
                task__in=user_tasks.filter(status='COMPLETED'),
                status='APPROVED'
            ).select_related('task')

            avg_completion_time = 0
            if completed_submissions.exists():
                times = []
                for sub in completed_submissions:
                    if sub.task.due_date and sub.submitted_at:
                        time_diff = (sub.submitted_at - sub.task.created_at).days
                        times.append(time_diff)
                if times:
                    avg_completion_time = sum(times) / len(times)

            productivity_score = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

            productivity_data.append({
                'user_id': user.id,
                'user_name': getattr(user, 'full_name', user.username),
                'tasks_completed': completed_tasks,
                'tasks_overdue': overdue_tasks,
                'average_completion_time': avg_completion_time,
                'productivity_score': productivity_score,
            })

        return productivity_data

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

class DepartmentAnalyticsView(generics.ListAPIView):
    serializer_class = DepartmentAnalyticsSerializer
    permission_classes = [IsAuthenticated, DepartmentAccessPermission]

    def get_queryset(self):
        departments = EmployeeProfile.objects.values_list('department', flat=True).distinct()

        analytics_data = []
        for dept in departments:
            if not dept:
                continue

            employees = EmployeeProfile.objects.filter(department=dept)
            total_employees = employees.count()

            tasks = Task.objects.filter(assigned_to__employee_profile__department=dept)
            total_tasks = tasks.count()
            completed_tasks = tasks.filter(status='COMPLETED').count()
            overdue_tasks = tasks.filter(status='OVERDUE').count()

            # Average productivity (simplified)
            avg_productivity = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

            analytics_data.append({
                'department': dept,
                'total_employees': total_employees,
                'total_tasks': total_tasks,
                'completed_tasks': completed_tasks,
                'overdue_tasks': overdue_tasks,
                'average_productivity': avg_productivity,
            })

        return analytics_data

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

class TaskTrendsView(generics.ListAPIView):
    serializer_class = TaskTrendsSerializer
    permission_classes = [IsAuthenticated, IsAdminOrManager]

    def get_queryset(self):
        # Trend analysis over the last 30 days
        end_date = now().date()
        start_date = end_date - timedelta(days=30)

        trends = Task.objects.filter(
            created_at__date__gte=start_date,
            created_at__date__lte=end_date
        ).annotate(
            date=TruncDate('created_at')
        ).values('date').annotate(
            tasks_created=Count('id'),
            tasks_completed=Count('id', filter=Q(status='COMPLETED')),
            tasks_overdue=Count('id', filter=Q(status='OVERDUE'))
        ).order_by('date')

        return list(trends)

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

class ExportReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrManager]

    def get(self, request, format_type):
        report_type = request.query_params.get('type', 'productivity')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')

        if format_type == 'excel':
            return self.export_excel(report_type, start_date, end_date)
        elif format_type == 'pdf':
            return self.export_pdf(report_type, start_date, end_date)
        else:
            return Response({'error': 'Invalid format'}, status=status.HTTP_400_BAD_REQUEST)

    def export_excel(self, report_type, start_date, end_date):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{report_type.capitalize()} Report"

        # Add headers and data based on report_type
        if report_type == 'productivity':
            ws.append(['User ID', 'User Name', 'Tasks Completed', 'Tasks Overdue', 'Avg Completion Time', 'Productivity Score'])
            view = EmployeeProductivityView()
            view.request = self.request
            data = view.get_queryset()
            for item in data:
                ws.append([
                    item['user_id'], item['user_name'], item['tasks_completed'],
                    item['tasks_overdue'], item['average_completion_time'], item['productivity_score']
                ])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = Response(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{report_type}_report.xlsx"'
        return response

    def export_pdf(self, report_type, start_date, end_date):
        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter

        p.drawString(100, height - 100, f"{report_type.capitalize()} Report")

        # Add data based on report_type
        y = height - 150
        if report_type == 'productivity':
            view = EmployeeProductivityView()
            view.request = self.request
            data = view.get_queryset()
            for item in data:
                p.drawString(100, y, f"{item['user_name']}: {item['productivity_score']}%")
                y -= 20
                if y < 100:
                    p.showPage()
                    y = height - 100

        p.showPage()
        p.save()
        buffer.seek(0)

        response = Response(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{report_type}_report.pdf"'
        return response

class ReportListCreateView(generics.ListCreateAPIView):
    queryset = Report.objects.all()
    serializer_class = ReportSerializer
    permission_classes = [IsAuthenticated, IsAdminOrManager]

    def perform_create(self, serializer):
        serializer.save(generated_by=self.request.user)


# Template-based Analytics Views

class EmployeeAnalyticsView(LoginRequiredMixin, TemplateView):
    template_name = 'analytics/employee_analytics.html'
    permission_classes = [EmployeeAnalyticsPermission]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        form = DateRangeForm(self.request.GET)
        context['form'] = form

        start_date = form.cleaned_data.get('start_date') if form.is_valid() else None
        end_date = form.cleaned_data.get('end_date') if form.is_valid() else None

        # Optimized queries for employee analytics
        tasks = Task.objects.filter(assigned_to=user).select_related('assigned_by').prefetch_related('time_logs', 'submissions')

        if start_date and end_date:
            tasks = tasks.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

        # KPIs
        total_tasks = tasks.count()
        completed_tasks = tasks.filter(status='APPROVED').count()
        overdue_tasks = tasks.filter(is_overdue=True).count()
        completion_rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

        # Time efficiency: total hours logged vs estimated (simplified)
        total_hours_logged = TimeLog.objects.filter(task__in=tasks, user=user).aggregate(total=Sum('hours'))['total'] or 0
        # Assuming estimated hours from task description or default
        estimated_hours = total_tasks * 8  # Assume 8 hours per task
        time_efficiency = (total_hours_logged / estimated_hours * 100) if estimated_hours > 0 else 0

        # Historical trends: monthly completion
        trends = tasks.annotate(month=TruncMonth('created_at')).values('month').annotate(
            completed=Count('id', filter=Q(status='APPROVED')),
            total=Count('id')
        ).order_by('month')

        context.update({
            'total_tasks': total_tasks,
            'completed_tasks': completed_tasks,
            'overdue_tasks': overdue_tasks,
            'completion_rate': completion_rate,
            'total_hours_logged': total_hours_logged,
            'time_efficiency': time_efficiency,
            'trends': list(trends),
        })
        return context


class ManagerAnalyticsView(LoginRequiredMixin, TemplateView):
    template_name = 'analytics/manager_analytics.html'
    permission_classes = [IsAdminOrManager]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        form = DateRangeForm(self.request.GET)
        context['form'] = form

        start_date = form.cleaned_data.get('start_date') if form.is_valid() else None
        end_date = form.cleaned_data.get('end_date') if form.is_valid() else None

        # Get manager's department
        try:
            manager = ManagerProfile.objects.get(user=user)
            department = manager.department
        except ManagerProfile.DoesNotExist:
            department = None

        # Employees under this manager
        employees = EmployeeProfile.objects.filter(manager=manager).select_related('user')

        # Tasks for employees in department
        tasks = Task.objects.filter(
            Q(assigned_to__employee_profile__manager=manager) |
            Q(assigned_to__employee_profile__department=department)
        ).select_related('assigned_by').prefetch_related('assigned_to', 'time_logs', 'submissions').distinct()

        if start_date and end_date:
            tasks = tasks.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

        # Department KPIs
        total_tasks = tasks.count()
        completed_tasks = tasks.filter(status='APPROVED').count()
        overdue_tasks = tasks.filter(is_overdue=True).count()
        completion_rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

        # Employee productivity
        employee_stats = []
        for emp in employees:
            emp_tasks = tasks.filter(assigned_to=emp.user)
            emp_completed = emp_tasks.filter(status='APPROVED').count()
            emp_total = emp_tasks.count()
            emp_rate = (emp_completed / emp_total * 100) if emp_total > 0 else 0
            employee_stats.append({
                'name': getattr(emp.user, 'full_name', emp.user.username),
                'total_tasks': emp_total,
                'completed_tasks': emp_completed,
                'completion_rate': emp_rate,
            })

        # Historical trends
        trends = tasks.annotate(month=TruncMonth('created_at')).values('month').annotate(
            completed=Count('id', filter=Q(status='APPROVED')),
            total=Count('id')
        ).order_by('month')

        context.update({
            'department': department,
            'total_tasks': total_tasks,
            'completed_tasks': completed_tasks,
            'overdue_tasks': overdue_tasks,
            'completion_rate': completion_rate,
            'employee_stats': employee_stats,
            'trends': list(trends),
        })
        return context


class DepartmentAnalyticsView(LoginRequiredMixin, TemplateView):
    template_name = 'analytics/department_analytics.html'
    permission_classes = [DepartmentAccessPermission]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        form = DateRangeForm(self.request.GET)
        context['form'] = form

        start_date = form.cleaned_data.get('start_date') if form.is_valid() else None
        end_date = form.cleaned_data.get('end_date') if form.is_valid() else None

        # Get user's department
        try:
            emp_profile = EmployeeProfile.objects.get(user=user)
            department = emp_profile.department
        except EmployeeProfile.DoesNotExist:
            department = None

        if not department and not user.is_staff:
            # If not employee and not admin, check manager
            try:
                manager = ManagerProfile.objects.get(user=user)
                department = manager.department
            except ManagerProfile.DoesNotExist:
                department = None

        # Tasks in department
        tasks = Task.objects.filter(
            assigned_to__employee_profile__department=department
        ).select_related('assigned_by').prefetch_related('assigned_to', 'time_logs', 'submissions').distinct()

        if start_date and end_date:
            tasks = tasks.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

        # Department KPIs
        total_tasks = tasks.count()
        completed_tasks = tasks.filter(status='APPROVED').count()
        overdue_tasks = tasks.filter(is_overdue=True).count()
        completion_rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

        # Employee productivity in department
        employees = EmployeeProfile.objects.filter(department=department).select_related('user')
        employee_stats = []
        for emp in employees:
            emp_tasks = tasks.filter(assigned_to=emp.user)
            emp_completed = emp_tasks.filter(status='APPROVED').count()
            emp_total = emp_tasks.count()
            emp_rate = (emp_completed / emp_total * 100) if emp_total > 0 else 0
            employee_stats.append({
                'name': getattr(emp.user, 'full_name', emp.user.username),
                'total_tasks': emp_total,
                'completed_tasks': emp_completed,
                'completion_rate': emp_rate,
            })

        # Historical trends
        trends = tasks.annotate(month=TruncMonth('created_at')).values('month').annotate(
            completed=Count('id', filter=Q(status='APPROVED')),
            total=Count('id')
        ).order_by('month')

        context.update({
            'department': department,
            'total_tasks': total_tasks,
            'completed_tasks': completed_tasks,
            'overdue_tasks': overdue_tasks,
            'completion_rate': completion_rate,
            'employee_stats': employee_stats,
            'trends': list(trends),
        })
        return context


class AdminAnalyticsView(LoginRequiredMixin, TemplateView):
    template_name = 'analytics/admin_analytics.html'
    permission_classes = [IsAdminOrManager]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        form = DateRangeForm(self.request.GET)
        context['form'] = form

        start_date = form.cleaned_data.get('start_date') if form.is_valid() else None
        end_date = form.cleaned_data.get('end_date') if form.is_valid() else None

        # All tasks
        tasks = Task.objects.all().select_related('assigned_by').prefetch_related('assigned_to', 'time_logs', 'submissions')

        if start_date and end_date:
            tasks = tasks.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

        # Overall KPIs
        total_tasks = tasks.count()
        completed_tasks = tasks.filter(status='APPROVED').count()
        overdue_tasks = tasks.filter(is_overdue=True).count()
        completion_rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0

        # Department-wise stats
        departments = EmployeeProfile.objects.values_list('department', flat=True).distinct()
        dept_stats = []
        for dept in departments:
            if not dept:
                continue
            dept_tasks = tasks.filter(assigned_to__employee_profile__department=dept)
            dept_completed = dept_tasks.filter(status='APPROVED').count()
            dept_total = dept_tasks.count()
            dept_rate = (dept_completed / dept_total * 100) if dept_total > 0 else 0
            dept_stats.append({
                'department': dept,
                'total_tasks': dept_total,
                'completed_tasks': dept_completed,
                'completion_rate': dept_rate,
            })

        # Historical trends
        trends = tasks.annotate(month=TruncMonth('created_at')).values('month').annotate(
            completed=Count('id', filter=Q(status='APPROVED')),
            total=Count('id')
        ).order_by('month')

        context.update({
            'total_tasks': total_tasks,
            'completed_tasks': completed_tasks,
            'overdue_tasks': overdue_tasks,
            'completion_rate': completion_rate,
            'dept_stats': dept_stats,
            'trends': list(trends),
        })
        return context


class TemplateExportView(LoginRequiredMixin, APIView):
    permission_classes = [IsAuthenticated, IsAdminOrManager]

    def get(self, request, format_type, view_type):
        start_date = request.GET.get('start_date')
        end_date = request.GET.get('end_date')

        # Get data based on view_type
        if view_type == 'employee':
            view = EmployeeAnalyticsView()
            view.request = request
            context = view.get_context_data()
        elif view_type == 'manager':
            view = ManagerAnalyticsView()
            view.request = request
            context = view.get_context_data()
        elif view_type == 'department':
            view = DepartmentAnalyticsView()
            view.request = request
            context = view.get_context_data()
        elif view_type == 'admin':
            view = AdminAnalyticsView()
            view.request = request
            context = view.get_context_data()
        else:
            return Response({'error': 'Invalid view type'}, status=status.HTTP_400_BAD_REQUEST)

        if format_type == 'excel':
            return self.export_excel(context, view_type)
        elif format_type == 'pdf':
            return self.export_pdf(context, view_type)
        else:
            return Response({'error': 'Invalid format'}, status=status.HTTP_400_BAD_REQUEST)

    def export_excel(self, context, view_type):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{view_type.capitalize()} Analytics Report"

        # Add headers and data based on context
        row = 1
        for key, value in context.items():
            if isinstance(value, list) and value:
                ws.cell(row=row, column=1, value=key.upper())
                row += 1
                if 'trends' in key:
                    ws.cell(row=row, column=1, value='Month')
                    ws.cell(row=row, column=2, value='Completed')
                    ws.cell(row=row, column=3, value='Total')
                    row += 1
                    for item in value:
                        ws.cell(row=row, column=1, value=str(item['month']))
                        ws.cell(row=row, column=2, value=item['completed'])
                        ws.cell(row=row, column=3, value=item['total'])
                        row += 1
                elif 'employee_stats' in key or 'dept_stats' in key:
                    ws.cell(row=row, column=1, value='Name/Department')
                    ws.cell(row=row, column=2, value='Total Tasks')
                    ws.cell(row=row, column=3, value='Completed Tasks')
                    ws.cell(row=row, column=4, value='Completion Rate')
                    row += 1
                    for item in value:
                        ws.cell(row=row, column=1, value=item.get('name', item.get('department', '')))
                        ws.cell(row=row, column=2, value=item['total_tasks'])
                        ws.cell(row=row, column=3, value=item['completed_tasks'])
                        ws.cell(row=row, column=4, value=item['completion_rate'])
                        row += 1
            elif not isinstance(value, (dict, list)):
                ws.cell(row=row, column=1, value=f"{key}: {value}")
                row += 1

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{view_type}_analytics.xlsx"'
        return response

    def export_pdf(self, context, view_type):
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()

        elements.append(Paragraph(f"{view_type.capitalize()} Analytics Report", styles['Title']))

        for key, value in context.items():
            if isinstance(value, list) and value:
                elements.append(Paragraph(key.upper(), styles['Heading2']))
                if 'trends' in key:
                    data = [['Month', 'Completed', 'Total']]
                    for item in value:
                        data.append([str(item['month']), str(item['completed']), str(item['total'])])
                    table = Table(data)
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black),
                    ]))
                    elements.append(table)
                elif 'employee_stats' in key or 'dept_stats' in key:
                    data = [['Name/Department', 'Total Tasks', 'Completed Tasks', 'Completion Rate']]
                    for item in value:
                        data.append([
                            item.get('name', item.get('department', '')),
                            str(item['total_tasks']),
                            str(item['completed_tasks']),
                            f"{item['completion_rate']:.2f}%"
                        ])
                    table = Table(data)
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black),
                    ]))
                    elements.append(table)
            elif not isinstance(value, (dict, list)):
                elements.append(Paragraph(f"{key}: {value}", styles['Normal']))

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{view_type}_analytics.pdf"'
        return response
